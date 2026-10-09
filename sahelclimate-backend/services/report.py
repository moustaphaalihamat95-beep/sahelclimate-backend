import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def generate_water_balance_excel(parcel_name: str, crop_type: str, records: list) -> io.BytesIO:
    """Génère un fichier Excel (.xlsx) stylisé du bilan hydrique en mémoire."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bilan Hydrique"

    # Style des en-têtes
    ws['A1'] = f"Rapport Bilan Hydrique - {parcel_name} ({crop_type})"
    ws['A1'].font = Font(size=14, bold=True, color="1B365D")
    
    headers = ["Date", "ET0 (mm/j)", "ETc (mm/j)", "Pluie (mm)", "Déficit (mm)", "Volume (m³)", "Statut"]
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_num in range(1, 8):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Données
    for r in records:
        ws.append([
            r.get("date", "N/A"),
            r.get("et0", 0.0),
            r.get("etc", 0.0),
            r.get("rain", 0.0),
            r.get("deficit", 0.0),
            r.get("volume_m3", 0.0),
            r.get("status", "OK")
        ])

    # Ajustement automatique des colonnes
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_water_balance_pdf(parcel_name: str, crop_type: str, records: list) -> io.BytesIO:
    """Génère un document PDF du bilan hydrique en mémoire."""
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#1B365D'), spaceAfter=12)
    meta_style = ParagraphStyle('MetaStyle', parent=styles['Normal'], fontSize=11, spaceAfter=20)

    story.append(Paragraph(f"🌾 <b>SahelClimate-AI — Rapport de Bilan Hydrique</b>", title_style))
    story.append(Paragraph(f"<b>Parcelle :</b> {parcel_name} | <b>Culture :</b> {crop_type}", meta_style))

    # En-têtes du tableau
    data = [["Date", "ET0 (mm)", "ETc (mm)", "Pluie (mm)", "Déficit", "Volume (m³)", "Statut"]]
    for r in records:
        data.append([
            r.get("date", "N/A"),
            f"{r.get('et0', 0.0):.2f}",
            f"{r.get('etc', 0.0):.2f}",
            f"{r.get('rain', 0.0):.1f}",
            f"{r.get('deficit', 0.0):.2f}",
            f"{r.get('volume_m3', 0.0):.1f}",
            r.get("status", "OK")
        ])

    t = Table(data, colWidths=[80, 65, 65, 65, 65, 75, 90])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')])
    ]))

    story.append(t)
    doc.build(story)
    output.seek(0)
    return output
